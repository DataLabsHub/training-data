#!/usr/bin/env python3
"""Self-contained mixed program and strict binary LLM verifier."""

from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import time
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

TESTS = Path(os.environ.get("SCF_TESTS_ROOT", Path(__file__).resolve().parent))
TASK_ROOT = TESTS.parent
DEFAULT_PUBLIC = Path("/app") if Path("/app").is_dir() else TASK_ROOT / "data" / "input_files"
PUBLIC_ROOT = Path(os.environ.get("SCF_PUBLIC_ROOT", DEFAULT_PUBLIC))
DEFAULT_CANDIDATE = Path("/app") if Path("/app").is_dir() else TESTS / "reference_answer"
CANDIDATE_ROOT = Path(os.environ.get("SCF_CANDIDATE_ROOT", DEFAULT_CANDIDATE))
LOGS = Path(os.environ.get("SCF_LOG_ROOT", "/logs/verifier"))

VERDICT_FIELDS = {
    "criterion_id",
    "criterion_type",
    "raw_score",
    "candidate_evidence",
    "reference_evidence",
    "reason",
    "material_contradiction",
}
VERDICT_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": sorted(VERDICT_FIELDS),
    "properties": {
        "criterion_id": {"type": "string"},
        "criterion_type": {"type": "string", "enum": ["positive", "penalty"]},
        "raw_score": {"type": "integer", "enum": [0, 1]},
        "candidate_evidence": {"type": "array", "items": {"type": "object"}},
        "reference_evidence": {"type": "array", "items": {"type": "object"}},
        "reason": {"type": "string"},
        "material_contradiction": {"type": "boolean"},
    },
}


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_relative(value: str) -> Path:
    pure = PurePosixPath(value)
    if pure.is_absolute() or ".." in pure.parts or pure.as_posix() != value:
        raise ValueError(f"unsafe relative path: {value}")
    return Path(*pure.parts)


def verify_public_inputs() -> dict[str, Any]:
    manifest = load_json(TESTS / "input_manifest.json")
    errors: list[str] = []
    for item in manifest["files"]:
        path = PUBLIC_ROOT / safe_relative(item["path"])
        if not path.is_file() or path.is_symlink():
            errors.append(f"missing_or_unsafe:{item['path']}")
        elif path.stat().st_size != item["size_bytes"] or sha256(path) != item["sha256"]:
            errors.append(f"modified:{item['path']}")
    return {
        "passed": not errors,
        "checked_files": len(manifest["files"]),
        "errors": errors,
    }


def run_program_check(criterion: dict[str, Any], candidate_root: Path) -> tuple[int, str]:
    params = criterion["method_params"]
    path = candidate_root / safe_relative(params["deliverable_path"])
    operator = params["operator"]
    try:
        if not path.is_file() or path.is_symlink():
            return 0, "declared deliverable is missing or unsafe"
        if operator == "file_exists":
            return 1, "regular file exists"
        if operator == "text_nonempty":
            return int(bool(path.read_text(encoding="utf-8").strip())), "UTF-8 text inspected"
        if operator == "json_parseable":
            load_json(path)
        elif operator == "csv_parseable":
            import csv

            with path.open("r", encoding="utf-8", newline="") as stream:
                next(csv.reader(stream))
        elif operator == "xlsx_openable":
            workbook = load_workbook(path, read_only=True, data_only=False)
            if not workbook.sheetnames:
                raise ValueError("workbook has no sheets")
            workbook.close()
        elif operator == "docx_openable":
            from docx import Document

            Document(path)
        elif operator == "pdf_openable":
            reader = PdfReader(path)
            if not reader.pages:
                raise ValueError("PDF has no pages")
        else:
            raise ValueError(f"unsupported program operator {operator}")
    except Exception as error:  # fail closed at the task boundary
        return 0, f"{type(error).__name__}: {error}"
    return 1, f"{operator} passed"


def validate_verdict(value: Any, criterion: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(value, dict) or set(value) != VERDICT_FIELDS:
        raise ValueError("Judge verdict fields do not match the strict schema")
    if value["criterion_id"] != criterion["id"]:
        raise ValueError("Judge criterion_id mismatch")
    if value["criterion_type"] != criterion["type"]:
        raise ValueError("Judge criterion_type mismatch")
    if type(value["raw_score"]) is not int or value["raw_score"] not in (0, 1):
        raise ValueError("Judge raw_score must be 0 or 1")
    if not isinstance(value["reason"], str) or not value["reason"].strip():
        raise ValueError("Judge reason is empty")
    if type(value["material_contradiction"]) is not bool:
        raise ValueError("material_contradiction must be boolean")
    for field in ("candidate_evidence", "reference_evidence"):
        if not isinstance(value[field], list):
            raise ValueError(f"{field} must be an array")
        for citation in value[field]:
            if not isinstance(citation, dict) or not isinstance(citation.get("path"), str):
                raise ValueError(f"{field} contains an invalid citation")
    candidate_roots = {
        item["path"].split("/", 1)[0] for item in value["candidate_evidence"]
    }
    reference_roots = {
        item["path"].split("/", 1)[0] for item in value["reference_evidence"]
    }
    if not candidate_roots <= {"candidate_answer", "public_input"}:
        raise ValueError("candidate evidence escaped the sealed evidence roots")
    if not reference_roots <= {"reference_answer", "public_input"}:
        raise ValueError("reference evidence escaped the sealed evidence roots")
    if value["raw_score"] == 1 and "candidate_answer" not in candidate_roots:
        raise ValueError("a positive pass or triggered penalty must cite candidate output")
    return value


def copy_declared_deliverables(source: Path, destination: Path) -> None:
    contract = load_json(TESTS / "deliverables.json")
    for item in contract["deliverables"]:
        source_path = source / safe_relative(item["path"])
        if not source_path.is_file() or source_path.is_symlink():
            continue
        target = destination / safe_relative(item["path"])
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_path, target)


def copy_public_inputs(destination: Path) -> None:
    manifest = load_json(TESTS / "input_manifest.json")
    for item in manifest["files"]:
        source = PUBLIC_ROOT / safe_relative(item["path"])
        target = destination / safe_relative(item["path"])
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)


def tree_hashes(root: Path) -> dict[str, str]:
    return {
        path.relative_to(root).as_posix(): sha256(path)
        for path in root.rglob("*")
        if path.is_file() and not path.is_symlink()
    }


def judge_prompt() -> str:
    return """You are a strict independent benchmark Judge. You are not the Executor.

Judge only the single criterion in criterion.json. Read judge_context.json for its exact frozen
business-truth slice and expert evidence. Use public_input/ for the original workspace,
candidate_answer/ for the submission, and reference_answer/ only as a comparison aid. Equality
with the reference is not enough; verify the candidate against the criterion and evidence.

For visual evidence, render and directly inspect the relevant PDF page, Office image, or raster.
Store any derivative only under _judge_scratch/. OCR, captions, filenames, and polished appearance
cannot replace pixel inspection.

For a positive criterion, raw_score=1 only when the submission completely satisfies the stated
condition. Any material omission, contradiction, wrong state/date/scope, or unavailable evidence
is raw_score=0. For a penalty criterion, raw_score=1 only when the exact prohibited overclaim is
materially present; otherwise raw_score=0. Never use partial credit and never return an unscored
state. Cite precise relative paths and locators. A score of 1 must cite candidate_answer/ itself.

Do not inspect paths outside this job and do not modify evidence. Write only the required JSON
object to verdict.json. Do not return chain-of-thought.
"""


def run_codex_judge(
    criterion: dict[str, Any],
    repeat: int,
    sealed_root: Path,
    evidence_digest: str,
) -> dict[str, Any]:
    model = os.environ.get("JUDGE_MODEL", "gpt-5.6-luna")
    effort = os.environ.get("JUDGE_REASONING_EFFORT", "high")
    identity = hashlib.sha256(
        json.dumps(
            {
                "evidence_digest": evidence_digest,
                "criterion": criterion,
                "repeat": repeat,
                "model": model,
                "reasoning_effort": effort,
            },
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()[:24]
    job = LOGS / "judge" / identity / criterion["id"] / f"repeat_{repeat:02d}"
    completed = job / "result.json"
    if completed.is_file():
        return validate_verdict(load_json(completed)["verdict"], criterion)
    job.mkdir(parents=True, exist_ok=True)
    for name in ("public_input", "candidate_answer", "reference_answer"):
        link = job / name
        if not link.exists():
            link.symlink_to(sealed_root / name, target_is_directory=True)
    (job / "_judge_scratch").mkdir(exist_ok=True)
    write_json(job / "criterion.json", criterion)
    shutil.copy2(TESTS / "judge_context" / f"{criterion['id']}.json", job / "judge_context.json")
    write_json(job / "verdict_schema.json", VERDICT_SCHEMA)
    maximum_attempts = int(os.environ.get("JUDGE_MAX_TRIES", "3"))
    timeout = int(os.environ.get("JUDGE_PER_TIMEOUT", "1200"))
    command = os.environ.get("JUDGE_CODEX_BIN", "codex")
    environment = os.environ.copy()
    if environment.get("JUDGE_API_KEY"):
        environment["OPENAI_API_KEY"] = environment["JUDGE_API_KEY"]
    if environment.get("JUDGE_BASE_URL"):
        environment["OPENAI_BASE_URL"] = environment["JUDGE_BASE_URL"]
    last_error = ""
    total_duration = 0.0
    for attempt in range(1, maximum_attempts + 1):
        verdict_path = job / "verdict.json"
        verdict_path.unlink(missing_ok=True)
        started = time.monotonic()
        process = subprocess.run(
            [
                command,
                "exec",
                "--model",
                model,
                "--config",
                f'model_reasoning_effort="{effort}"',
                "--dangerously-bypass-approvals-and-sandbox",
                "--skip-git-repo-check",
                "--ephemeral",
                "--ignore-rules",
                "--json",
                "--cd",
                str(job),
                "-",
            ],
            input=judge_prompt(),
            text=True,
            capture_output=True,
            timeout=timeout,
            check=False,
            env=environment,
        )
        duration = time.monotonic() - started
        total_duration += duration
        (job / f"attempt_{attempt:02d}.trace.jsonl").write_text(
            process.stdout, encoding="utf-8"
        )
        (job / f"attempt_{attempt:02d}.stderr.log").write_text(
            process.stderr, encoding="utf-8"
        )
        try:
            if process.returncode != 0:
                raise RuntimeError(f"Codex exited {process.returncode}")
            verdict = validate_verdict(load_json(verdict_path), criterion)
        except Exception as error:
            last_error = f"{type(error).__name__}: {error}"
            continue
        write_json(
            completed,
            {
                "verdict": verdict,
                "attempts": attempt,
                "duration_seconds": total_duration,
                "judge_model": model,
            },
        )
        return verdict
    raise RuntimeError(last_error or "Judge exhausted retries")


def aggregate_score(
    criteria: list[dict[str, Any]], raw_scores: dict[str, int]
) -> float:
    denominator = sum(int(item["weight"]) for item in criteria if item["type"] == "positive")
    positive = sum(
        int(item["weight"]) * raw_scores[item["id"]]
        for item in criteria
        if item["type"] == "positive"
    )
    penalty = sum(
        abs(int(item["weight"])) * raw_scores[item["id"]]
        for item in criteria
        if item["type"] == "penalty"
    )
    return max(0.0, min(1.0, (positive - penalty) / denominator))


def validate_only(criteria: list[dict[str, Any]]) -> int:
    input_report = verify_public_inputs()
    program = {
        item["id"]: run_program_check(item, CANDIDATE_ROOT)
        for item in criteria
        if item["method"] == "program"
    }
    missing_context = [
        item["id"]
        for item in criteria
        if item["method"] == "llm_judge"
        and not (TESTS / "judge_context" / f"{item['id']}.json").is_file()
    ]
    report = {
        "status": (
            "valid"
            if input_report["passed"]
            and not missing_context
            and all(reward == 1 for reward, _reason in program.values())
            else "invalid"
        ),
        "mode": "validate_only_no_llm_score",
        "public_inputs": input_report,
        "program_checks": {
            key: {"raw_score": value[0], "reason": value[1]}
            for key, value in program.items()
        },
        "missing_judge_context": missing_context,
        "criterion_count": len(criteria),
    }
    write_json(LOGS / "validation.json", report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "valid" else 1


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    LOGS.mkdir(parents=True, exist_ok=True)
    rubric = load_json(TESTS / "rubric.json")
    criteria: list[dict[str, Any]] = rubric["criteria"]
    if args.validate_only:
        return validate_only(criteria)
    input_report = verify_public_inputs()
    if not input_report["passed"]:
        write_json(
            LOGS / "reward.json",
            {"reward": 0.0, "status": "input_integrity_failed", "details": input_report},
        )
        return 0
    backend = os.environ.get("JUDGE_BACKEND", "codex")
    if backend != "codex":
        write_json(
            LOGS / "reward.json",
            {
                "reward": 0.0,
                "status": "unsupported_judge_backend",
                "error": "Production scoring requires JUDGE_BACKEND=codex",
            },
        )
        return 0
    program_results = {
        item["id"]: run_program_check(item, CANDIDATE_ROOT)
        for item in criteria
        if item["method"] == "program"
    }
    repeats = int(os.environ.get("JUDGE_REPEATS", "1"))
    if repeats < 1 or repeats % 2 == 0:
        raise ValueError("JUDGE_REPEATS must be a positive odd number")
    concurrency = max(1, int(os.environ.get("JUDGE_CONCURRENCY", "4")))
    semantic = [item for item in criteria if item["method"] == "llm_judge"]
    with tempfile.TemporaryDirectory(prefix="sealed-judge-evidence-") as raw:
        sealed = Path(raw)
        copy_public_inputs(sealed / "public_input")
        copy_declared_deliverables(CANDIDATE_ROOT, sealed / "candidate_answer")
        shutil.copytree(TESTS / "reference_answer", sealed / "reference_answer")
        frozen_hashes = tree_hashes(sealed)
        evidence_digest = hashlib.sha256(
            json.dumps(frozen_hashes, sort_keys=True).encode("utf-8")
        ).hexdigest()
        with concurrent.futures.ThreadPoolExecutor(max_workers=concurrency) as pool:
            futures = {
                pool.submit(
                    run_codex_judge,
                    criterion,
                    repeat,
                    sealed,
                    evidence_digest,
                ): (
                    criterion["id"],
                    repeat,
                )
                for criterion in semantic
                for repeat in range(1, repeats + 1)
            }
            verdicts: dict[str, list[dict[str, Any]]] = {
                item["id"]: [] for item in semantic
            }
            errors: list[str] = []
            for future, (criterion_id, repeat) in futures.items():
                try:
                    verdicts[criterion_id].append(future.result())
                except Exception as error:
                    errors.append(
                        f"{criterion_id}/repeat_{repeat:02d}:"
                        f"{type(error).__name__}:{error}"
                    )
        if tree_hashes(sealed) != frozen_hashes:
            errors.append("Judge modified sealed evidence")
    if errors:
        write_json(
            LOGS / "reward.json",
            {"reward": 0.0, "status": "judge_failed_closed", "errors": errors},
        )
        return 0
    raw_scores = {key: value[0] for key, value in program_results.items()}
    for criterion_id, values in verdicts.items():
        rewards = [int(item["raw_score"]) for item in values]
        raw_scores[criterion_id] = int(sum(rewards) > len(rewards) // 2)
    reward = aggregate_score(criteria, raw_scores)
    write_json(
        LOGS / "reward.json",
        {
            "reward": reward,
            "status": "scored",
            "pass_threshold": rubric["pass_threshold"],
            "passed": reward >= float(rubric["pass_threshold"]),
            "criterion_scores": raw_scores,
            "program_results": {
                key: {"raw_score": value[0], "reason": value[1]}
                for key, value in program_results.items()
            },
            "judge_repeats": repeats,
            "judge_model": os.environ.get("JUDGE_MODEL", "gpt-5.6-luna"),
        },
    )
    (LOGS / "reward.txt").write_text(f"{reward:.10f}\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
